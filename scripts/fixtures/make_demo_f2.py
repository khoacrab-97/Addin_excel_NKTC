"""Tạo workbook demo cho F2 từ bảng tiến độ thật (sheet TDTC của test_addin_xuat_nktc.xlsx).

Sheet tạo ra:
- TDTC: bảng tiến độ gốc.
- NK_TheoNgay: kết quả F1 (giữ tất cả các dòng, quyết định Q7) + 5 cột phụ với dữ liệu MẪU.
- Form: biểu mẫu nhật ký thi công MẪU (ô merge, 1 ô công thức) — chưa phải biểu mẫu thật của dự án.

Cách dùng:
    python3 -I scripts/fixtures/make_demo_f2.py test/fixtures/test_addin_xuat_nktc.xlsx test/fixtures/demo_f2.xlsx
"""
import datetime
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side

EPOCH = datetime.datetime(1899, 12, 30)
WEATHER = ["Nắng", "Nắng", "Nhiều mây", "Mưa rào chiều", "Nắng"]
INSPECTIONS = {
    "30/06/2026": "Nghiệm thu công tác chuẩn bị mặt bằng, lán trại",
    "17/08/2026": "Nghiệm thu lắp đặt đường ống tưới, thử áp lực đường ống",
    "31/08/2026": "Nghiệm thu lắp đặt thiết bị tưới tự động",
    "30/12/2026": "Nghiệm thu trồng cây xanh, bó bồn gốc cây",
}


def serial(value):
    return (value - EPOCH).days


def main():
    source, output = sys.argv[1], sys.argv[2]
    tasks_wb = openpyxl.load_workbook(source)
    tasks = tasks_wb["TDTC"]
    rows = [(str(r[0]).strip(), serial(r[1]), serial(r[2])) for r in tasks.iter_rows(min_row=2, values_only=True) if r[0]]

    wb = openpyxl.Workbook()
    copy = wb.active
    copy.title = "TDTC"
    for row in tasks.iter_rows(values_only=True):
        copy.append(list(row))
    for cell in copy["B"][1:] + copy["C"][1:]:
        cell.number_format = "dd/mm/yyyy"

    daily = wb.create_sheet("NK_TheoNgay")
    daily.append(["Ngày", "Nội dung", "Thời tiết", "Ban chỉ huy - CBKT", "Số công nhân", "Số thiết bị", "Nội dung nghiệm thu"])
    first, last = min(r[1] for r in rows), max(r[2] for r in rows)
    for offset, day in enumerate(range(first, last + 1)):
        content = "\n".join(f"- {name}" for name, start, end in rows if start <= day <= end)
        active = sum(1 for _, start, end in rows if start <= day <= end)
        label = (EPOCH + datetime.timedelta(days=day)).strftime("%d/%m/%Y")
        daily.append([
            EPOCH + datetime.timedelta(days=day),
            content,
            WEATHER[offset % len(WEATHER)],
            "Nguyễn Văn A (CHT); Trần Văn B (CBKT)",
            8 + active * 2,
            2 + active // 3,
            INSPECTIONS.get(label, ""),
        ])
    for cell in daily["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for cell in daily["B"][1:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    daily.column_dimensions["B"].width = 80

    form = wb.create_sheet("Form")
    thin = Side(style="thin")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, width in zip("ABCDEF", [20, 14, 14, 14, 14, 14]):
        form.column_dimensions[col].width = width
    def merge(rng, value=None, bold=False, center=False, wrap=False):
        form.merge_cells(rng)
        cell = form[rng.split(":")[0]]
        if value is not None:
            cell.value = value
        cell.font = Font(bold=bold)
        cell.alignment = Alignment(horizontal="center" if center else None, vertical="top", wrap_text=wrap)
    merge("A1:F1", "NHẬT KÝ THI CÔNG", bold=True, center=True)
    merge("A2:F2", "Công trình: Cảnh quan – Hạ tầng (biểu mẫu MẪU để chạy thử)", center=True)
    form["A4"] = "Ngày:"
    merge("B4:C4")
    form["B4"].number_format = "dd/mm/yyyy"
    form["D4"] = "Thời tiết:"
    merge("E4:F4")
    form["A5"] = "Ban chỉ huy - CBKT:"
    merge("B5:F5")
    form["A6"] = "Số công nhân:"
    form["D6"] = "Số thiết bị:"
    merge("A8:F8", "1. NỘI DUNG CÔNG VIỆC THỰC HIỆN", bold=True)
    merge("A9:F18", wrap=True)
    merge("A20:F20", "2. NỘI DUNG NGHIỆM THU", bold=True)
    merge("A21:F24", wrap=True)
    merge("A26:F26", '="Xác nhận nhật ký ngày "&TEXT(B4,"dd/mm/yyyy")')
    for row in form["A4:F24"]:
        for cell in row:
            cell.border = box
    wb.save(output)
    print(f"Đã ghi {output}: {last - first + 1} ngày")


if __name__ == "__main__":
    main()
