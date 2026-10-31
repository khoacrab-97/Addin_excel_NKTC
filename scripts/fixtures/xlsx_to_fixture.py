"""Chuyển một sheet .xlsx thành fixture JSON mô phỏng `range.values` của Office.js.

- Ngày (datetime) → Excel serial date (số), giống giá trị Office.js trả về.
- Ô trống → "" (Office.js trả chuỗi rỗng).
- Công thức → giá trị đã tính lưu trong file (data_only).

Cách dùng (chạy với -I để cô lập môi trường Python):
    python3 -I scripts/fixtures/xlsx_to_fixture.py <file.xlsx> <output.json> [tên sheet]
    python3 -I scripts/fixtures/xlsx_to_fixture.py <file.xlsx> <output.json> --all   # mọi sheet + công thức
"""
import datetime
import json
import sys

import openpyxl
from openpyxl.utils import get_column_letter

EPOCH_1900 = datetime.datetime(1899, 12, 30)


def to_office_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime.datetime):
        delta = value - EPOCH_1900
        return delta.days + delta.seconds / 86400
    if isinstance(value, datetime.date):
        return (datetime.datetime(value.year, value.month, value.day) - EPOCH_1900).days
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value
    return str(value)


def sheet_fixture(sheet, formula_sheet):
    values = [[to_office_value(cell.value) for cell in row] for row in sheet.iter_rows()]
    formulas = {
        cell.coordinate: cell.value
        for row in formula_sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, str) and cell.value.startswith("=")
    }
    return {
        "name": sheet.title,
        "address": f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}",
        "values": values,
        "formulas": formulas,
    }


def main_all(source, output):
    values_wb = openpyxl.load_workbook(source, data_only=True)
    formulas_wb = openpyxl.load_workbook(source)
    fixture = {"sheets": [sheet_fixture(s, formulas_wb[s.title]) for s in values_wb.worksheets]}
    with open(output, "w", encoding="utf-8") as handle:
        json.dump(fixture, handle, ensure_ascii=False, indent=1)
    print(f"Đã ghi {output}: {', '.join(s['name'] for s in fixture['sheets'])}")


def main():
    source, output = sys.argv[1], sys.argv[2]
    if len(sys.argv) > 3 and sys.argv[3] == "--all":
        main_all(source, output)
        return
    workbook = openpyxl.load_workbook(source, data_only=True)
    if workbook.epoch != openpyxl.utils.datetime.CALENDAR_WINDOWS_1900:
        raise SystemExit("Chỉ hỗ trợ workbook hệ ngày 1900.")
    sheet = workbook[sys.argv[3]] if len(sys.argv) > 3 else workbook.worksheets[0]
    values = [[to_office_value(cell.value) for cell in row] for row in sheet.iter_rows()]
    fixture = {
        "sheetName": sheet.title,
        "address": f"A1:{get_column_letter(sheet.max_column)}{sheet.max_row}",
        "values": values,
    }
    with open(output, "w", encoding="utf-8") as handle:
        json.dump(fixture, handle, ensure_ascii=False, indent=1)
    print(f"Đã ghi {output}: sheet '{sheet.title}', {sheet.max_row} dòng x {sheet.max_column} cột")


if __name__ == "__main__":
    main()
