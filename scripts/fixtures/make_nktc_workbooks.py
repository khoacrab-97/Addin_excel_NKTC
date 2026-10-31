"""Tạo biểu mẫu Excel "Tình hình thi công hàng ngày" (chuyển từ mẫu Word ND_Nhật ký thi công)
và workbook demo cho F2.

Kết quả:
- templates/Mau_NKTC.xlsx: biểu mẫu trống (sheet "Mẫu NKTC") để dùng thực tế.
- test/fixtures/demo_f2.xlsx: TDTC (tiến độ thật) + NK_TheoNgay (kết quả F1 + cột phụ MẪU) + "Mẫu NKTC".
- test/fixtures/demo_f2.targets.json: ô đích của từng cột nguồn (dùng cho smoke test F2).

Cách dùng:
    python3 -I scripts/fixtures/make_nktc_workbooks.py test/fixtures/test_addin_xuat_nktc.xlsx
"""
import datetime
import json
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break

FORM_SHEET = "Mẫu NKTC"
FONT = "Times New Roman"
COLUMNS = 9  # A..I
COLUMN_WIDTH = 9.6  # ~16,5 cm vùng in (A4, lề trái 2,5 cm, phải 2 cm); in vừa 1 trang ngang
LINE_HEIGHT = 15.75
CM_PER_INCH = 2.54
EPOCH = datetime.datetime(1899, 12, 30)
LAST = get_column_letter(COLUMNS)

DOTTED = Side(style="dotted")
DATE_FORMAT = '"Ngày "dd" tháng "mm" năm "yyyy'


class FormBuilder:
    """Dựng biểu mẫu theo từng dòng; ghi lại ô đích cho từng trường dữ liệu."""

    def __init__(self, sheet):
        self.ws = sheet
        self.row = 1
        self.targets = {}

    def font(self, cell, bold=False, italic=False, size=12):
        cell.font = Font(name=FONT, bold=bold, italic=italic, size=size)

    def text(self, value, bold=False, italic=False, size=12, indent=0, height_lines=1, align=None):
        """Dòng chữ chiếm toàn bộ chiều ngang (merge A:I)."""
        rng = f"A{self.row}:{LAST}{self.row + height_lines - 1}"
        self.ws.merge_cells(rng)
        cell = self.ws[f"A{self.row}"]
        cell.value = value
        self.font(cell, bold, italic, size)
        cell.alignment = Alignment(horizontal=align, vertical="top", wrap_text=True, indent=indent)
        for r in range(self.row, self.row + height_lines):
            self.ws.row_dimensions[r].height = LINE_HEIGHT
        self.row += height_lines

    def block(self, field, lines):
        """Vùng nhập nhiều dòng (merge A:I), gạch chấm phía dưới như mẫu Word."""
        top, bottom = self.row, self.row + lines - 1
        self.ws.merge_cells(f"A{top}:{LAST}{bottom}")
        cell = self.ws[f"A{top}"]
        self.font(cell)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        for r in range(top, bottom + 1):
            self.ws.row_dimensions[r].height = LINE_HEIGHT
            for c in range(1, COLUMNS + 1):
                self.ws.cell(r, c).border = Border(bottom=DOTTED)
        self.targets[field] = f"A{top}"
        self.row = bottom + 1

    def handwriting(self, lines):
        """Dòng gạch chấm để ghi tay (không phải ô đích)."""
        for _ in range(lines):
            for c in range(1, COLUMNS + 1):
                self.ws.cell(self.row, c).border = Border(bottom=DOTTED)
            self.ws.row_dimensions[self.row].height = LINE_HEIGHT * 1.3
            self.row += 1

    def labeled(self, label, field, first_col=3, last_col=COLUMNS, unit=None, number_format=None):
        """Nhãn ở A:B, ô đích ở các cột còn lại (merge), có thể kèm đơn vị ở cột cuối."""
        r = self.row
        self.ws.merge_cells(f"A{r}:B{r}")
        self.ws[f"A{r}"] = label
        self.font(self.ws[f"A{r}"])
        value_last = last_col - (1 if unit else 0)
        rng = f"{get_column_letter(first_col)}{r}:{get_column_letter(value_last)}{r}"
        if first_col != value_last:
            self.ws.merge_cells(rng)
        cell = self.ws[f"{get_column_letter(first_col)}{r}"]
        self.font(cell)
        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if number_format:
            cell.number_format = number_format
        for c in range(first_col, value_last + 1):
            self.ws.cell(r, c).border = Border(bottom=DOTTED)
        if unit:
            self.ws.cell(r, last_col).value = unit
            self.font(self.ws.cell(r, last_col))
        self.ws.row_dimensions[r].height = LINE_HEIGHT
        self.targets[field] = f"{get_column_letter(first_col)}{r}"
        self.row += 1


def build_form(ws):
    for c in range(1, COLUMNS + 1):
        ws.column_dimensions[get_column_letter(c)].width = COLUMN_WIDTH
    f = FormBuilder(ws)

    f.text("TÌNH HÌNH THI CÔNG HÀNG NGÀY", bold=True, size=13, align="center")
    # Ngày: ô đích nhận Excel serial date; định dạng số hiển thị "Ngày dd tháng mm năm yyyy".
    r = f.row
    ws.merge_cells(f"E{r}:{LAST}{r}")
    date_cell = ws[f"E{r}"]
    date_cell.number_format = DATE_FORMAT
    f.font(date_cell, italic=True)
    date_cell.alignment = Alignment(horizontal="right")
    ws.row_dimensions[r].height = LINE_HEIGHT
    f.targets["Ngày"] = f"E{r}"
    f.row += 2

    f.text("1. Thời tiết", bold=True)
    checkbox_row = f.row
    f.row += 1
    f.labeled("Ghi chú thời tiết:", "Thời tiết")
    weather = f"$C${f.row - 1}"
    has = lambda word: f'ISNUMBER(SEARCH("{word}",{weather}))'  # noqa: E731
    checkboxes = [
        ("B", f'=IF(AND({weather}<>"",NOT({has("mưa")}),NOT({has("nắng")})),"☒","☐")&" Bình thường"'),
        ("E", f'=IF({has("mưa")},"☒","☐")&" Mưa"'),
        ("H", f'=IF({has("nắng")},"☒","☐")&" Nắng"'),
    ]
    for col, formula in checkboxes:
        cell = ws[f"{col}{checkbox_row}"]
        cell.value = formula
        f.font(cell)
    ws.row_dimensions[checkbox_row].height = LINE_HEIGHT
    r = f.row
    ws[f"A{r}"] = "Nhiệt độ:"
    f.font(ws[f"A{r}"])
    for label_col, value_col, label, field in [
        ("B", "C", "Sáng", "Nhiệt độ sáng"),
        ("D", "E", "Trưa", "Nhiệt độ trưa"),
        ("F", "G", "Chiều", "Nhiệt độ chiều"),
        ("H", "I", "Tối", "Nhiệt độ tối"),
    ]:
        ws[f"{label_col}{r}"] = label
        f.font(ws[f"{label_col}{r}"])
        ws[f"{label_col}{r}"].alignment = Alignment(horizontal="right")
        value = ws[f"{value_col}{r}"]
        f.font(value)
        value.alignment = Alignment(horizontal="center")
        value.number_format = '0"°C"'
        value.border = Border(bottom=DOTTED)
        f.targets[field] = f"{value_col}{r}"
    ws.row_dimensions[r].height = LINE_HEIGHT
    f.row += 2

    f.text("2. Tình hình nhân lực, thiết bị thi công:", bold=True)
    f.text("2.1. Nhân lực:", bold=True, indent=2)
    f.labeled("Ban chỉ huy - CBKT:", "Ban chỉ huy - CBKT")
    f.labeled("Số lượng công nhân:", "Số công nhân", last_col=5, unit="người")
    f.block("Nhân lực - chi tiết", 2)
    f.text("2.2. Thiết bị:", bold=True, indent=2)
    f.labeled("Số lượng thiết bị:", "Số thiết bị", last_col=5, unit="chiếc")
    f.block("Thiết bị - chi tiết", 2)

    f.text("3. Tình hình thi công:", bold=True)
    f.text("3.1. Nội dung, khối lượng công việc và tiến độ thực hiện:", bold=True, indent=2)
    f.block("Nội dung", 18)
    f.text("3.2. Nghiệm thu công việc xây dựng hàng ngày:", bold=True, indent=2)
    f.block("Nội dung nghiệm thu", 8)

    # Ngắt trang trước mục 4: in 2 trang A4 giống mẫu Word.
    ws.row_breaks.append(Break(id=f.row - 1))
    f.text(
        "4. Mô tả chi tiết các sự cố, hư hỏng và các vấn đề phát sinh khác trong quá trình thi công xây dựng công trình:",
        bold=True,
        height_lines=2,
    )
    f.block("Sự cố, phát sinh", 4)
    f.text(
        "5. Các kiến nghị và những ý kiến chỉ đạo giải quyết các vấn đề phát sinh của các bên có liên quan "
        "(Chủ đầu tư, Người giám sát thi công, Người giám sát tác giả)",
        bold=True,
        height_lines=2,
    )
    f.block("Kiến nghị", 4)
    f.text(
        "6. Nhận xét của kỹ sư giám sát thi công xây dựng về tình hình và chất lượng công việc xây dựng "
        "(ghi rõ họ và tên, chức vụ và chữ ký)",
        bold=True,
        height_lines=2,
    )
    f.handwriting(4)
    f.text("7. Chỉ huy trưởng công trường tiếp thu các kiến nghị và nhận xét (họ tên, chữ ký)", bold=True)
    f.handwriting(4)

    # In: A4 đứng, lề như mẫu Word, vừa 1 trang theo chiều ngang.
    ws.print_area = f"A1:{LAST}{f.row - 1}"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins = PageMargins(
        left=2.5 / CM_PER_INCH, right=2.0 / CM_PER_INCH, top=2.0 / CM_PER_INCH, bottom=2.0 / CM_PER_INCH
    )
    ws.sheet_view.showGridLines = False
    return f.targets


WEATHER = ["Nắng", "Nắng", "Nhiều mây", "Nắng, chiều mưa rào", "Bình thường"]
INSPECTIONS = {
    "30/06/2026": "Nghiệm thu công tác chuẩn bị mặt bằng, lán trại",
    "17/08/2026": "Nghiệm thu lắp đặt đường ống tưới, thử áp lực đường ống",
    "31/08/2026": "Nghiệm thu lắp đặt thiết bị tưới tự động",
    "30/12/2026": "Nghiệm thu trồng cây xanh, bó bồn gốc cây",
}
INCIDENTS = {"05/08/2026": "Mưa lớn chiều, tạm dừng đào mương cáp 2 giờ."}
REQUESTS = {"05/08/2026": "Đề nghị TVGS kiểm tra lại cao độ đáy mương cáp đoạn băng đường."}
DAILY_HEADERS = [
    "Ngày",
    "Nội dung",
    "Thời tiết",
    "Ban chỉ huy - CBKT",
    "Số công nhân",
    "Số thiết bị",
    "Nội dung nghiệm thu",
    "Sự cố, phát sinh",
    "Kiến nghị",
]


def serial(value):
    return (value - EPOCH).days


def build_demo(source, form_targets):
    tasks = openpyxl.load_workbook(source)["TDTC"]
    rows = [(str(r[0]).strip(), serial(r[1]), serial(r[2])) for r in tasks.iter_rows(min_row=2, values_only=True) if r[0]]

    wb = openpyxl.Workbook()
    copy = wb.active
    copy.title = "TDTC"
    for row in tasks.iter_rows(values_only=True):
        copy.append(list(row))
    for cell in copy["B"][1:] + copy["C"][1:]:
        cell.number_format = "dd/mm/yyyy"

    daily = wb.create_sheet("NK_TheoNgay")
    daily.append(DAILY_HEADERS)
    first, last = min(r[1] for r in rows), max(r[2] for r in rows)
    for offset, day in enumerate(range(first, last + 1)):
        active = [name for name, start, end in rows if start <= day <= end]
        label = (EPOCH + datetime.timedelta(days=day)).strftime("%d/%m/%Y")
        daily.append([
            EPOCH + datetime.timedelta(days=day),
            "\n".join(f"- {name}" for name in active),
            WEATHER[offset % len(WEATHER)],
            "Nguyễn Văn A (CHT); Trần Văn B (CBKT)",
            8 + len(active) * 2,
            2 + len(active) // 3,
            INSPECTIONS.get(label, ""),
            INCIDENTS.get(label, ""),
            REQUESTS.get(label, ""),
        ])
    for cell in daily["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for cell in daily["B"][1:]:
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    daily.column_dimensions["B"].width = 80

    build_form(wb.create_sheet(FORM_SHEET))
    targets = [
        {"column": get_column_letter(i + 1), "header": header, "cell": form_targets[header]}
        for i, header in enumerate(DAILY_HEADERS)
    ]
    return wb, last - first + 1, targets


def main():
    source = sys.argv[1]
    template = openpyxl.Workbook()
    targets = build_form(template.active)
    template.active.title = FORM_SHEET
    template.save("templates/Mau_NKTC.xlsx")

    demo, days, mapping = build_demo(source, targets)
    demo.save("test/fixtures/demo_f2.xlsx")
    form = demo[FORM_SHEET]
    merged = {str(rng).split(":")[0]: str(rng) for rng in form.merged_cells.ranges}
    for entry in mapping:
        entry["select"] = merged.get(entry["cell"], entry["cell"])
    meta = {"sheet": FORM_SHEET, "mappings": mapping}
    meta["formulaCell"] = next(
        f"{c.column_letter}{c.row}" for row in form.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")
    )
    with open("test/fixtures/demo_f2.targets.json", "w", encoding="utf-8") as handle:
        json.dump(meta, handle, ensure_ascii=False, indent=1)
    print("Đã ghi templates/Mau_NKTC.xlsx, test/fixtures/demo_f2.xlsx", f"({days} ngày), test/fixtures/demo_f2.targets.json")
    for entry in mapping:
        print(f"  {entry['column']} {entry['header']:<22} → {FORM_SHEET}!{entry['cell']} (chọn {entry['select']})")


if __name__ == "__main__":
    main()
